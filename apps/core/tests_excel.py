"""Pruebas del Excel del catálogo (Sección 10, punto 9).

Se corren con: python manage.py test apps.core.tests_excel
(usan una base de datos de prueba; no tocan el catálogo real).
"""

from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.core.excel import (
    ANTIVIBRATION_PREFIX,
    PRODUCTS_SHEET,
    build_catalog_workbook,
    import_catalog_workbook,
)
from apps.core.models import AntivibrationTechnicalData, CatalogItem


def edit_workbook(buffer, edit):
    """Abre el Excel exportado, aplica `edit(sheet, headers)` y lo devuelve como archivo."""
    workbook = load_workbook(buffer)
    sheet = workbook[PRODUCTS_SHEET]
    headers = [cell.value for cell in sheet[1]]
    edit(sheet, headers)
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def column(headers, name):
    return headers.index(name) + 1


class CatalogExcelTests(TestCase):

    def setUp(self):
        self.cdc = CatalogItem.objects.create(
            name="CDC-2",
            category="antivibratorios",
            subcategory="colgantes",
            description="<p>Colgante</p>",
        )
        AntivibrationTechnicalData.objects.create(
            product=self.cdc,
            model_code="CDC-2",
            capacity_kg=Decimal("45.00"),
            elastomer_material="Neopreno",
        )
        self.ap = CatalogItem.objects.create(
            name="AP-151",
            category="accionamiento",
        )

    def test_export_has_all_products_and_labels(self):
        workbook = load_workbook(build_catalog_workbook())
        sheet = workbook[PRODUCTS_SHEET]
        headers = [cell.value for cell in sheet[1]]
        rows = list(sheet.iter_rows(min_row=2, values_only=True))

        self.assertEqual(len(rows), 2)
        self.assertIn("Instrucciones", workbook.sheetnames)
        cdc = next(row for row in rows if row[1] == "CDC-2")
        self.assertEqual(cdc[column(headers, "Categoría") - 1], "Antivibratorios")
        self.assertEqual(cdc[column(headers, "Subcategoría") - 1], "Colgantes antivibración")
        self.assertEqual(cdc[column(headers, ANTIVIBRATION_PREFIX + "Capacidad (kg)") - 1], 45.0)

    def test_round_trip_without_changes(self):
        result = import_catalog_workbook(build_catalog_workbook())

        self.assertTrue(result.ok, result.errors)
        self.assertEqual((result.created, result.updated, result.unchanged), (0, 0, 2))
        self.cdc.refresh_from_db()
        self.assertEqual(self.cdc.subcategory, "colgantes")
        self.assertEqual(self.cdc.antivibration_data.capacity_kg, Decimal("45.00"))

    def test_edit_capacity_and_add_new_product(self):
        def edit(sheet, headers):
            for row in sheet.iter_rows(min_row=2):
                if row[1].value == "CDC-2":
                    row[column(headers, ANTIVIBRATION_PREFIX + "Capacidad (kg)") - 1].value = 60
            new_row = [None] * len(headers)
            new_row[column(headers, "Clave") - 1] = "CDC-3"
            new_row[column(headers, "Categoría") - 1] = "antivibratorios"
            new_row[column(headers, "Subcategoría") - 1] = "Colgantes antivibración"
            new_row[column(headers, ANTIVIBRATION_PREFIX + "Capacidad (kg)") - 1] = "80"
            sheet.append(new_row)

        result = import_catalog_workbook(edit_workbook(build_catalog_workbook(), edit))

        self.assertTrue(result.ok, result.errors)
        self.assertEqual((result.created, result.updated, result.unchanged), (1, 1, 1))
        self.cdc.refresh_from_db()
        self.assertEqual(self.cdc.antivibration_data.capacity_kg, Decimal("60.00"))
        new = CatalogItem.objects.get(name="CDC-3")
        self.assertEqual(new.antivibration_data.model_code, "CDC-3")
        self.assertEqual(new.antivibration_data.capacity_kg, Decimal("80.00"))

    def test_empty_technical_cell_keeps_value(self):
        def edit(sheet, headers):
            for row in sheet.iter_rows(min_row=2):
                if row[1].value == "CDC-2":
                    row[column(headers, ANTIVIBRATION_PREFIX + "Material del elastómero") - 1].value = None
                    row[column(headers, ANTIVIBRATION_PREFIX + "Capacidad (kg)") - 1].value = 50

        result = import_catalog_workbook(edit_workbook(build_catalog_workbook(), edit))

        self.assertTrue(result.ok, result.errors)
        data = AntivibrationTechnicalData.objects.get(product=self.cdc)
        self.assertEqual(data.elastomer_material, "Neopreno")
        self.assertEqual(data.capacity_kg, Decimal("50.00"))

    def test_errors_save_nothing(self):
        def edit(sheet, headers):
            for row in sheet.iter_rows(min_row=2):
                if row[1].value == "CDC-2":
                    row[column(headers, "Activo") - 1].value = "No"
                if row[1].value == "AP-151":
                    row[column(headers, "Categoría") - 1].value = "Inventada"

        result = import_catalog_workbook(edit_workbook(build_catalog_workbook(), edit))

        self.assertFalse(result.ok)
        self.assertIn("Inventada", " ".join(result.errors))
        self.cdc.refresh_from_db()
        self.assertTrue(self.cdc.is_active)

    def test_duplicate_key_is_rejected(self):
        def edit(sheet, headers):
            new_row = [None] * len(headers)
            new_row[column(headers, "Clave") - 1] = "cdc-2"
            new_row[column(headers, "Categoría") - 1] = "Antivibratorios"
            sheet.append(new_row)

        result = import_catalog_workbook(edit_workbook(build_catalog_workbook(), edit))

        self.assertFalse(result.ok)
        self.assertEqual(CatalogItem.objects.count(), 2)

    def test_template_has_headers_only(self):
        sheet = load_workbook(build_catalog_workbook(template=True))[PRODUCTS_SHEET]
        self.assertEqual(sheet.max_row, 1)


class CatalogExcelViewTests(TestCase):

    def setUp(self):
        CatalogItem.objects.create(name="AP-151", category="accionamiento")
        self.user = get_user_model().objects.create_user("marcos", password="x")
        for codename in ("view_catalogitem", "add_catalogitem", "change_catalogitem"):
            self.user.user_permissions.add(Permission.objects.get(codename=codename))

    def test_anonymous_cannot_download(self):
        response = self.client.get(reverse("exportar_excel"))
        self.assertEqual(response.status_code, 302)

    def test_download_and_upload(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse("exportar_excel"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])

        upload = SimpleUploadedFile("catalogo.xlsx", response.content)
        response = self.client.post(reverse("importar_excel"), {"archivo": upload}, follow=True)

        messages = [str(message) for message in response.context["messages"]]
        self.assertTrue(any("1 sin cambios" in message for message in messages), messages)

    def test_rejects_non_excel(self):
        self.client.force_login(self.user)
        upload = SimpleUploadedFile("datos.csv", b"a,b")
        response = self.client.post(reverse("importar_excel"), {"archivo": upload}, follow=True)
        messages = [str(message) for message in response.context["messages"]]
        self.assertTrue(any(".xlsx" in message for message in messages), messages)
