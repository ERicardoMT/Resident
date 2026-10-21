"""Exportar e importar el catálogo en Excel (Sección 10, punto 9 del libro maestro).

- Exportar: un libro con la hoja "Productos" (todas las columnas, incluidos
  los datos técnicos), "Instrucciones" y, si hay permiso, "Mediciones".
- Importar: el mismo formato. Cada renglón crea o actualiza un producto por
  ID o, si el ID está vacío, por clave (nombre). Si algún renglón tiene
  error no se guarda nada.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import (
    SUBCATEGORIES_BY_CATEGORY,
    AntivibrationTechnicalData,
    CatalogCategory,
    CatalogItem,
    CatalogSubcategory,
    CustomCatalogCategory,
    CustomCatalogSubcategory,
    LevelerTechnicalData,
)


PRODUCTS_SHEET = "Productos"
INSTRUCTIONS_SHEET = "Instrucciones"
MEASUREMENTS_SHEET = "Mediciones"

MAX_IMPORT_ROWS = 2000
MAX_ERRORS_SHOWN = 20

# Columnas del producto: (encabezado, campo del modelo)
PRODUCT_COLUMNS = [
    ("ID", "id"),
    ("Clave", "name"),
    ("Categoría", "category"),
    ("Subcategoría", "subcategory"),
    ("Descripción (HTML)", "description"),
    ("Precio", "price_label"),
    ("Etiqueta", "badge"),
    ("Activo", "is_active"),
    ("Orden", "sort_order"),
]

# Archivos: se exportan como referencia; no se importan desde Excel
FILE_COLUMNS = [
    ("Imagen (solo lectura)", "image"),
    ("Modelo 3D (solo lectura)", "model_3d"),
    ("Modelo AR (solo lectura)", "ar_model"),
]

ANTIVIBRATION_PREFIX = "Antivibratorio · "
ANTIVIBRATION_COLUMNS = [
    ("Código de modelo", "model_code"),
    ("Diámetro de base", "base_diameter"),
    ("Altura de base", "base_height"),
    ("Diámetro de tornillo", "screw_diameter"),
    ("Altura de tornillo", "screw_height"),
    ("Capacidad (texto)", "capacity_label"),
    ("Capacidad (kg)", "capacity_kg"),
    ("Material del elastómero", "elastomer_material"),
    ("Material del tornillo", "screw_material"),
]

LEVELER_PREFIX = "Nivelador · "
LEVELER_COLUMNS = [
    ("Código de modelo", "model_code"),
    ("Capacidad (texto)", "capacity_label"),
    ("Capacidad (kg)", "capacity_kg"),
    ("Tipo", "type_label"),
    ("Roscas métricas", "metric_threads"),
    ("Roscas estándar", "standard_threads"),
    ("Diámetro de base", "base_diameter"),
    ("Altura de tornillo", "screw_height"),
    ("Material del tornillo", "screw_material"),
    ("Material de la base", "base_material"),
    ("URL del producto", "product_url"),
]

MEASUREMENT_COLUMNS = [
    ("Folio", "folio"),
    ("Fecha", "created_at"),
    ("Frecuencia dominante (Hz)", "dominant_hz"),
    ("RPM", "rpm"),
    ("Muestreo (Hz)", "sample_rate_hz"),
    ("Muestras", "sample_count"),
    ("Duración (s)", "duration_s"),
    ("RMS (m/s²)", "rms_ms2"),
    ("Pico (m/s²)", "peak_ms2"),
    ("Archivo PDF", "pdf_filename"),
]

TRUE_WORDS = {"si", "sí", "s", "yes", "y", "1", "true", "verdadero", "activo"}
FALSE_WORDS = {"no", "n", "0", "false", "falso", "inactivo"}

HEADER_FILL = PatternFill("solid", fgColor="063B68")
HEADER_FONT = Font(bold=True, color="FFFFFF")
READONLY_FILL = PatternFill("solid", fgColor="E6E9EE")


def _headers() -> list[str]:
    return (
        [header for header, _ in PRODUCT_COLUMNS]
        + [header for header, _ in FILE_COLUMNS]
        + [ANTIVIBRATION_PREFIX + header for header, _ in ANTIVIBRATION_COLUMNS]
        + [LEVELER_PREFIX + header for header, _ in LEVELER_COLUMNS]
    )


# =====================================================
# Categorías: se exportan con su nombre visible y se
# aceptan por nombre o por código al importar
# =====================================================

def _category_maps() -> tuple[dict[str, str], dict[str, str]]:
    """Devuelve (texto -> código, código -> nombre visible)."""
    to_code: dict[str, str] = {}
    to_label: dict[str, str] = {}

    for value, label in CatalogCategory.choices:
        to_code[value.casefold()] = value
        to_code[label.casefold()] = value
        to_label[value] = label

    for name in CustomCatalogCategory.objects.values_list("name", flat=True):
        to_code.setdefault(name.casefold(), name)
        to_label.setdefault(name, name)

    return to_code, to_label


def _subcategory_maps() -> tuple[dict[str, str], dict[str, str], dict[str, set[str]]]:
    """Devuelve (texto -> código, código -> nombre, categoría -> subcategorías válidas)."""
    to_code: dict[str, str] = {}
    to_label: dict[str, str] = {}

    for value, label in CatalogSubcategory.choices:
        to_code[value.casefold()] = value
        to_code[label.casefold()] = value
        to_label[value] = label

    allowed: dict[str, set[str]] = {
        str(category): {str(sub) for sub in subcategories}
        for category, subcategories in SUBCATEGORIES_BY_CATEGORY.items()
    }

    for category, name in CustomCatalogSubcategory.objects.values_list(
        "category",
        "name",
    ):
        to_code.setdefault(name.casefold(), name)
        to_label.setdefault(name, name)
        allowed.setdefault(category, set()).add(name)

    return to_code, to_label, allowed


# =====================================================
# Exportar
# =====================================================

def _cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, Decimal):
        return float(value)
    return value


def _style_header(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "C2"
    sheet.row_dimensions[1].height = 32


def _write_products(sheet, items) -> None:
    category_labels = _category_maps()[1]
    subcategory_labels = _subcategory_maps()[1]

    sheet.append(_headers())

    for item in items:
        av = getattr(item, "antivibration_data", None)
        lv = getattr(item, "leveler_data", None)

        row = [
            item.id,
            item.name,
            category_labels.get(item.category, item.category),
            subcategory_labels.get(item.subcategory or "", item.subcategory or ""),
            item.description,
            item.price_label,
            item.badge,
            "Sí" if item.is_active else "No",
            item.sort_order,
        ]
        row += [
            getattr(item, field_name).name if getattr(item, field_name) else ""
            for _, field_name in FILE_COLUMNS
        ]
        row += [
            _cell_value(getattr(av, field_name)) if av else ""
            for _, field_name in ANTIVIBRATION_COLUMNS
        ]
        row += [
            _cell_value(getattr(lv, field_name)) if lv else ""
            for _, field_name in LEVELER_COLUMNS
        ]
        sheet.append(row)

    widths = [8, 18, 22, 26, 40, 12, 12, 9, 8] + [22] * len(FILE_COLUMNS)
    widths += [20] * (len(ANTIVIBRATION_COLUMNS) + len(LEVELER_COLUMNS))
    _style_header(sheet, widths)

    # Columnas de solo lectura en gris
    first_file_col = len(PRODUCT_COLUMNS) + 1
    for column in range(first_file_col, first_file_col + len(FILE_COLUMNS)):
        sheet.cell(row=1, column=column).fill = READONLY_FILL
        sheet.cell(row=1, column=column).font = Font(bold=True, color="5F6B7A")


def _write_instructions(sheet) -> None:
    category_labels = _category_maps()[1]
    subcategory_labels = _subcategory_maps()[1]
    allowed = _subcategory_maps()[2]

    lines = [
        ["Cómo usar este Excel"],
        [""],
        ["1. Edita la hoja \"Productos\" y súbela desde el panel con \"Subir Excel\"."],
        ["2. Si el ID tiene número, se actualiza ese producto (puedes cambiarle la clave)."],
        ["3. Si el ID está vacío, se busca la clave: si existe se actualiza; si no, se crea un producto nuevo."],
        ["4. Clave y Categoría son obligatorias. Activo acepta Sí o No."],
        ["5. Las columnas \"solo lectura\" (imagen y modelos 3D) no se importan: esos archivos se suben desde Editar producto."],
        ["6. Los datos técnicos solo se guardan si llenas al menos una columna de su grupo (Antivibratorio · o Nivelador ·). Una celda vacía no borra el dato."],
        ["7. Si algún renglón tiene error no se guarda nada, y el panel te dice qué renglón corregir."],
        ["8. No cambies los nombres de las columnas ni el nombre de la hoja \"Productos\"."],
        [""],
        ["Categorías válidas", "Subcategorías válidas"],
    ]
    for code, label in category_labels.items():
        subs = sorted(
            subcategory_labels.get(sub, sub)
            for sub in allowed.get(code, set())
        )
        lines.append([label, ", ".join(subs) if subs else "(sin subcategoría)"])

    for line in lines:
        sheet.append(line)

    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A12"].font = Font(bold=True)
    sheet["B12"].font = Font(bold=True)
    sheet.column_dimensions["A"].width = 40
    sheet.column_dimensions["B"].width = 90


def _write_measurements(sheet) -> None:
    from apps.vibration.models import VibrationMeasurement

    sheet.append([header for header, _ in MEASUREMENT_COLUMNS])

    # Sin el PDF ni el análisis completo: no van en el Excel
    for measurement in VibrationMeasurement.objects.defer("pdf_data", "analysis_data").order_by("-created_at"):
        row = []
        for _, field_name in MEASUREMENT_COLUMNS:
            value = getattr(measurement, field_name)
            if field_name == "created_at" and value:
                value = timezone.localtime(value).strftime("%Y-%m-%d %H:%M:%S")
            row.append(_cell_value(value))
        sheet.append(row)

    _style_header(sheet, [24, 20] + [16] * (len(MEASUREMENT_COLUMNS) - 2))


def build_catalog_workbook(*, include_measurements: bool = False, template: bool = False) -> BytesIO:
    """Arma el Excel del catálogo. Con template=True deja solo encabezados e instrucciones."""
    workbook = Workbook()
    products = workbook.active
    products.title = PRODUCTS_SHEET

    items = (
        []
        if template
        else CatalogItem.objects.select_related(
            "antivibration_data",
            "leveler_data",
        ).order_by("category", "sort_order", "name")
    )
    _write_products(products, items)

    _write_instructions(workbook.create_sheet(INSTRUCTIONS_SHEET))

    if include_measurements and not template:
        _write_measurements(workbook.create_sheet(MEASUREMENTS_SHEET))

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer


# =====================================================
# Importar
# =====================================================

@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    errors: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors


def _same(current: Any, new: Any) -> bool:
    """Igualdad tolerante: Excel cambia \\r\\n por \\n y vacío por None."""
    def normalize(value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, Decimal):
            return str(value.quantize(Decimal("0.01")))
        return str(value).replace("\r\n", "\n").strip()
    return normalize(current) == normalize(new)


def _apply(instance, values: dict[str, Any]) -> bool:
    """Asigna solo los campos que cambiaron; devuelve True si hubo cambios."""
    changed = False
    for field_name, value in values.items():
        if not _same(getattr(instance, field_name), value):
            setattr(instance, field_name, value)
            changed = True
    return changed


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def _parse_bool(value: Any, row_number: int, errors: list[str]) -> bool | None:
    text = _text(value).casefold()
    if not text:
        return None
    if text in TRUE_WORDS:
        return True
    if text in FALSE_WORDS:
        return False
    errors.append(f"Renglón {row_number}: \"Activo\" debe ser Sí o No (dice \"{_text(value)}\").")
    return None


def _parse_int(value: Any, label: str, row_number: int, errors: list[str]) -> int | None:
    text = _text(value)
    if not text:
        return None
    try:
        number = int(float(text))
    except ValueError:
        errors.append(f"Renglón {row_number}: \"{label}\" debe ser un número entero (dice \"{text}\").")
        return None
    if number < 0:
        errors.append(f"Renglón {row_number}: \"{label}\" no puede ser negativo.")
        return None
    return number


def _parse_decimal(value: Any, label: str, row_number: int, errors: list[str]) -> Decimal | None:
    text = _text(value).replace(",", "")
    if not text:
        return None
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation:
        errors.append(f"Renglón {row_number}: \"{label}\" debe ser un número (dice \"{_text(value)}\").")
        return None


def _technical_values(row: dict[str, Any], prefix: str, columns, row_number: int, errors: list[str]) -> dict[str, Any] | None:
    """Valores del grupo técnico; None si el renglón no trae ninguno."""
    raw = {field_name: row.get(prefix + header) for header, field_name in columns}
    if not any(_text(value) for value in raw.values()):
        return None

    # Solo las celdas con valor: una celda vacía no borra el dato guardado
    values: dict[str, Any] = {}
    for header, field_name in columns:
        if not _text(raw[field_name]):
            continue
        if field_name == "capacity_kg":
            values[field_name] = _parse_decimal(raw[field_name], prefix + header, row_number, errors)
        else:
            values[field_name] = _text(raw[field_name])
    return values


def import_catalog_workbook(uploaded_file) -> ImportResult:
    """Valida todo el archivo y, solo si no hay errores, guarda los cambios."""
    result = ImportResult()

    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except Exception:
        result.errors.append("No se pudo leer el archivo. Súbelo en formato .xlsx.")
        return result

    if PRODUCTS_SHEET not in workbook.sheetnames:
        result.errors.append(f"El archivo no tiene la hoja \"{PRODUCTS_SHEET}\". Usa el Excel que descargas del panel.")
        return result

    rows = list(workbook[PRODUCTS_SHEET].iter_rows(values_only=True))
    if not rows:
        result.errors.append("La hoja \"Productos\" está vacía.")
        return result

    headers = [_text(header) for header in rows[0]]
    missing = [header for header in ("ID", "Clave", "Categoría") if header not in headers]
    if missing:
        result.errors.append("Faltan las columnas: " + ", ".join(missing) + ". No cambies los encabezados.")
        return result

    data_rows = [
        (index, dict(zip(headers, values)))
        for index, values in enumerate(rows[1:], start=2)
        if any(_text(value) for value in values)
    ]
    if len(data_rows) > MAX_IMPORT_ROWS:
        result.errors.append(f"El archivo tiene más de {MAX_IMPORT_ROWS} productos.")
        return result

    category_codes = _category_maps()[0]
    subcategory_codes, _, allowed_subcategories = _subcategory_maps()

    items_by_id = {item.id: item for item in CatalogItem.objects.all()}
    ids_by_name = {item.name.casefold(): item.id for item in items_by_id.values()}
    av_owner = dict(AntivibrationTechnicalData.objects.values_list("model_code", "product_id"))
    lv_owner = dict(LevelerTechnicalData.objects.values_list("model_code", "product_id"))
    av_code_by_product = {product_id: code for code, product_id in av_owner.items()}
    lv_code_by_product = {product_id: code for code, product_id in lv_owner.items()}

    errors = result.errors
    planned = []
    seen_names: dict[str, int] = {}
    seen_ids: dict[int, int] = {}
    seen_codes: dict[str, int] = {}

    for row_number, row in data_rows:
        name = _text(row.get("Clave"))
        if not name:
            errors.append(f"Renglón {row_number}: falta la Clave.")
            continue
        if len(name) > 160:
            errors.append(f"Renglón {row_number}: la Clave tiene más de 160 caracteres.")
            continue
        if name.casefold() in seen_names:
            errors.append(f"Renglón {row_number}: la clave \"{name}\" ya aparece en el renglón {seen_names[name.casefold()]}.")
            continue
        seen_names[name.casefold()] = row_number

        # Producto existente: por ID o, si no hay ID, por clave
        item_id = _parse_int(row.get("ID"), "ID", row_number, errors)
        item = None
        if item_id is not None:
            item = items_by_id.get(item_id)
            if item is None:
                errors.append(f"Renglón {row_number}: el ID {item_id} no existe. Déjalo vacío para crear un producto nuevo.")
                continue
            if item_id in seen_ids:
                errors.append(f"Renglón {row_number}: el ID {item_id} ya aparece en el renglón {seen_ids[item_id]}.")
                continue
            seen_ids[item_id] = row_number
            other_id = ids_by_name.get(name.casefold())
            if other_id is not None and other_id != item_id:
                errors.append(f"Renglón {row_number}: ya existe otro producto con la clave \"{name}\" (ID {other_id}).")
                continue
        else:
            existing_id = ids_by_name.get(name.casefold())
            if existing_id is not None:
                if existing_id in seen_ids:
                    errors.append(f"Renglón {row_number}: el producto \"{name}\" ya se editó en el renglón {seen_ids[existing_id]}.")
                    continue
                seen_ids[existing_id] = row_number
                item = items_by_id[existing_id]

        category_text = _text(row.get("Categoría"))
        category = category_codes.get(category_text.casefold())
        if not category:
            errors.append(f"Renglón {row_number}: la categoría \"{category_text}\" no existe. Revisa la hoja Instrucciones.")
            continue

        subcategory_text = _text(row.get("Subcategoría"))
        subcategory = ""
        if subcategory_text:
            subcategory = subcategory_codes.get(subcategory_text.casefold(), "")
            if subcategory not in allowed_subcategories.get(category, set()):
                errors.append(f"Renglón {row_number}: la subcategoría \"{subcategory_text}\" no corresponde a la categoría \"{category_text}\".")
                continue

        values = {
            "name": name,
            "category": category,
            "subcategory": subcategory or None,
        }
        for header, field_name in (("Descripción (HTML)", "description"), ("Precio", "price_label"), ("Etiqueta", "badge")):
            if header in headers:
                values[field_name] = _text(row.get(header))
        if values.get("price_label") == "":
            values["price_label"] = "Cotizar"
        if len(values.get("price_label", "")) > 40 or len(values.get("badge", "")) > 40:
            errors.append(f"Renglón {row_number}: \"Precio\" y \"Etiqueta\" admiten hasta 40 caracteres.")
            continue

        is_active = _parse_bool(row.get("Activo"), row_number, errors)
        if is_active is not None:
            values["is_active"] = is_active
        sort_order = _parse_int(row.get("Orden"), "Orden", row_number, errors)
        if sort_order is not None:
            values["sort_order"] = sort_order

        av_values = _technical_values(row, ANTIVIBRATION_PREFIX, ANTIVIBRATION_COLUMNS, row_number, errors)
        lv_values = _technical_values(row, LEVELER_PREFIX, LEVELER_COLUMNS, row_number, errors)

        # Código de modelo: obligatorio y único en cada grupo técnico
        for group_values, owners, code_by_product, label in (
            (av_values, av_owner, av_code_by_product, ANTIVIBRATION_PREFIX),
            (lv_values, lv_owner, lv_code_by_product, LEVELER_PREFIX),
        ):
            if group_values is None:
                continue
            # Sin código en el Excel: se conserva el actual o se usa la clave
            code = (
                group_values.get("model_code")
                or (code_by_product.get(item.id) if item else None)
                or name
            )
            group_values["model_code"] = code
            owner = owners.get(code)
            if owner is not None and (item is None or owner != item.id):
                errors.append(f"Renglón {row_number}: \"{label}Código de modelo\" {code} ya pertenece a otro producto.")
            seen_key = label + code.casefold()
            if seen_key in seen_codes:
                errors.append(f"Renglón {row_number}: \"{label}Código de modelo\" {code} ya aparece en el renglón {seen_codes[seen_key]}.")
            seen_codes[seen_key] = row_number

        planned.append((item, values, av_values, lv_values))

    if errors:
        return result

    with transaction.atomic():
        for item, values, av_values, lv_values in planned:
            is_new = item is None
            if is_new:
                item = CatalogItem.objects.create(**values)
                changed = True
            else:
                changed = _apply(item, values)
                if changed:
                    item.save()

            for model, group_values in (
                (AntivibrationTechnicalData, av_values),
                (LevelerTechnicalData, lv_values),
            ):
                if group_values is None:
                    continue
                technical = model.objects.filter(product=item).first()
                if technical is None:
                    model.objects.create(product=item, **group_values)
                    changed = True
                elif _apply(technical, group_values):
                    technical.save()
                    changed = True

            if is_new:
                result.created += 1
            elif changed:
                result.updated += 1
            else:
                result.unchanged += 1

    return result
