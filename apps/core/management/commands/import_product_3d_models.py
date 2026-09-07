from __future__ import annotations

import csv
import hashlib
import re
import struct
import unicodedata

from pathlib import Path, PureWindowsPath
from urllib.parse import unquote, urlparse
from urllib.request import Request, urlopen

from django.core.files.base import ContentFile
from django.core.files.storage import FileSystemStorage
from django.core.management.base import BaseCommand, CommandError
from django.utils.text import slugify

from openpyxl import load_workbook

from apps.core.models import (
    AntivibrationTechnicalData,
    CatalogItem,
)


ERROR_STATUSES = {
    "SIN_3D",
    "SIN_PRODUCTO",
    "PRODUCTO_AMBIGUO",
    "DUPLICADO_EXCEL",
    "FUENTE_INVALIDA",
    "GLB_INVALIDO",
}


READY_STATUSES = {
    "LISTO",
    "LISTO_REEMPLAZO",
}


def exact_key(value):
    """
    Clave estricta para códigos de modelo.

    Conserva caracteres significativos como:
    *, -, espacios, etc.

    Ejemplos:
    MSF-3*2 != MSF-32
    MPC-3*2 != MPC-32
    """
    return str(
        value or ""
    ).strip().casefold()


def normalize(value):
    """
    Clave flexible usada solamente como respaldo.

    Nunca se utiliza antes de intentar una
    coincidencia exacta.
    """
    text = str(
        value or ""
    ).strip()

    text = unicodedata.normalize(
        "NFKD",
        text,
    )

    text = "".join(
        char
        for char in text
        if not unicodedata.combining(char)
    )

    return re.sub(
        r"[^a-z0-9]+",
        "",
        text.lower(),
    )


def extract_source(cell):
    """
    Obtiene el origen real del 3D.

    Prioridad:
    1. Hipervínculo de Excel.
    2. Fórmula HYPERLINK.
    3. Valor visible de la celda.
    """

    if (
        cell.hyperlink
        and cell.hyperlink.target
    ):
        return str(
            cell.hyperlink.target
        ).strip()

    value = cell.value

    if value is None:
        return ""

    value = str(value).strip()

    match = re.match(
        r'^\s*=HYPERLINK\(\s*"([^"]+)"',
        value,
        flags=re.IGNORECASE,
    )

    if match:
        return match.group(1).strip()

    return value


def is_url(value):
    parsed = urlparse(
        str(value or "")
    )

    return parsed.scheme.lower() in {
        "http",
        "https",
    }


def filename_from_source(source):
    source = str(
        source or ""
    ).strip()

    if not source:
        return ""

    if is_url(source):
        return Path(
            unquote(
                urlparse(source).path
            )
        ).name

    return PureWindowsPath(
        source
    ).name


def validate_glb(data):
    """
    Valida estructura mínima de un GLB 2.0.

    Header GLB:
    - magic
    - version
    - total length
    """

    if len(data) < 12:
        raise ValueError(
            "El archivo es demasiado pequeño "
            "para ser un GLB."
        )

    magic, version, declared_length = (
        struct.unpack(
            "<III",
            data[:12],
        )
    )

    if magic != 0x46546C67:
        raise ValueError(
            "El archivo descargado no tiene "
            "cabecera GLB/glTF. "
            "Puede ser una página web en lugar "
            "del archivo 3D."
        )

    if version != 2:
        raise ValueError(
            f"GLB versión {version}. "
            "Se requiere glTF 2.0."
        )

    if declared_length != len(data):
        raise ValueError(
            "La longitud declarada por el GLB "
            "no coincide con el archivo real."
        )

    return {
        "version": version,
        "bytes": len(data),
    }


class Command(BaseCommand):

    help = (
        "Importa los modelos GLB de los "
        "antivibratorios usando Modelos3D.xlsx."
    )


    def add_arguments(
        self,
        parser,
    ):
        parser.add_argument(
            "--excel",
            required=True,
        )

        parser.add_argument(
            "--models-dir",
            default="",
        )

        parser.add_argument(
            "--sheet",
            default="Antivibratorios",
        )

        parser.add_argument(
            "--dry-run",
            action="store_true",
        )

        parser.add_argument(
            "--apply",
            action="store_true",
        )

        parser.add_argument(
            "--replace-existing",
            action="store_true",
        )

        parser.add_argument(
            "--allow-filesystem-storage",
            action="store_true",
        )

        parser.add_argument(
            "--max-file-mb",
            type=int,
            default=250,
        )


    def handle(
        self,
        *args,
        **options,
    ):
        if (
            options["dry_run"]
            and options["apply"]
        ):
            raise CommandError(
                "No combines --dry-run y --apply."
            )

        apply_changes = bool(
            options["apply"]
        )

        excel = Path(
            options["excel"]
        ).expanduser()

        if not excel.exists():
            raise CommandError(
                f"No existe el Excel: {excel}"
            )

        excel = excel.resolve()


        models_dir = None

        if options["models_dir"]:
            models_dir = Path(
                options["models_dir"]
            ).expanduser()

            if not models_dir.exists():
                raise CommandError(
                    "No existe la carpeta de "
                    f"modelos: {models_dir}"
                )

            models_dir = (
                models_dir.resolve()
            )


        workbook = load_workbook(
            excel,
            data_only=False,
            read_only=False,
        )


        sheet_name = options["sheet"]

        if sheet_name not in workbook.sheetnames:
            raise CommandError(
                f"No existe la hoja "
                f"'{sheet_name}'. "
                f"Hojas: "
                f"{', '.join(workbook.sheetnames)}"
            )


        worksheet = workbook[
            sheet_name
        ]


        headers = {}

        for column in range(
            1,
            worksheet.max_column + 1,
        ):
            value = worksheet.cell(
                1,
                column,
            ).value

            if value is not None:
                headers[
                    str(value).strip()
                ] = column


        if "Modelo" not in headers:
            raise CommandError(
                "No encontré la columna 'Modelo'."
            )


        if "3D" not in headers:
            raise CommandError(
                "No encontré la columna '3D'."
            )


        model_column = headers[
            "Modelo"
        ]

        source_column = headers[
            "3D"
        ]


        self.stdout.write("")
        self.stdout.write("=" * 78)
        self.stdout.write(
            "SMAV INAHER - AUDITORÍA DE MODELOS 3D"
        )
        self.stdout.write("=" * 78)
        self.stdout.write(
            f"Excel        : {excel}"
        )
        self.stdout.write(
            f"Hoja         : {sheet_name}"
        )
        self.stdout.write(
            f"Filas        : "
            f"{worksheet.max_row - 1}"
        )
        self.stdout.write(
            f"Columna modelo: {model_column}"
        )
        self.stdout.write(
            f"Columna 3D    : {source_column}"
        )
        self.stdout.write("")


        technical_records = list(
            AntivibrationTechnicalData
            .objects
            .select_related("product")
            .order_by("model_code")
        )


        exact_products_by_model = {}
        normalized_products_by_model = {}

        for technical in technical_records:

            exact = exact_key(
                technical.model_code
            )

            normalized = normalize(
                technical.model_code
            )

            exact_products_by_model.setdefault(
                exact,
                [],
            ).append(
                technical.product
            )

            normalized_products_by_model.setdefault(
                normalized,
                [],
            ).append(
                technical.product
            )


        exact_fallback_by_name = {}
        normalized_fallback_by_name = {}

        fallback_products = (
            CatalogItem.objects
            .filter(
                category="antivibratorios",
                is_active=True,
            )
            .order_by("id")
        )

        for product in fallback_products:

            exact = exact_key(
                product.name
            )

            normalized = normalize(
                product.name
            )

            exact_fallback_by_name.setdefault(
                exact,
                [],
            ).append(product)

            normalized_fallback_by_name.setdefault(
                normalized,
                [],
            ).append(product)


        field = (
            CatalogItem
            ._meta
            .get_field("model_3d")
        )

        storage = field.storage

        storage_name = (
            storage.__class__.__module__
            + "."
            + storage.__class__.__name__
        )


        self.stdout.write(
            f"Storage      : {storage_name}"
        )
        self.stdout.write("")


        if (
            apply_changes
            and isinstance(
                storage,
                FileSystemStorage,
            )
            and not options[
                "allow_filesystem_storage"
            ]
        ):
            raise CommandError(
                "\nIMPORTACIÓN BLOQUEADA.\n\n"
                "El proyecto está usando "
                "FileSystemStorage.\n"
                "Ejecutar --apply desde esta PC "
                "guardaría los GLB en el disco local "
                "de esta computadora.\n\n"
                "Para la carga real deberás ejecutar "
                "el comando en el servidor con MEDIA "
                "persistente, o configurar storage remoto."
            )


        max_bytes = (
            options["max_file_mb"]
            * 1024
            * 1024
        )


        rows = []


        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):
            model = str(
                worksheet.cell(
                    row_number,
                    model_column,
                ).value
                or ""
            ).strip()

            source = extract_source(
                worksheet.cell(
                    row_number,
                    source_column,
                )
            )

            if not model:
                continue


            result = {
                "excel_row": row_number,
                "model": model,
                "source": source,
                "product_id": "",
                "product_name": "",
                "resolved_source": "",
                "status": "",
                "detail": "",
                "bytes": "",
                "sha256": "",
                "stored_name": "",
            }


            if not source:
                result["status"] = (
                    "SIN_3D"
                )

                result["detail"] = (
                    "La celda 3D está vacía."
                )

                rows.append(result)
                continue


            exact = exact_key(
                model
            )

            normalized = normalize(
                model
            )


            # Primero intentamos el código exacto.
            #
            # MSF-3*2 debe coincidir únicamente con
            # MSF-3*2 y NO con MSF-32.
            matches = (
                exact_products_by_model.get(
                    exact,
                    [],
                )
            )


            if not matches:
                matches = (
                    exact_fallback_by_name.get(
                        exact,
                        [],
                    )
                )


            # Solamente cuando no existe coincidencia
            # exacta usamos la comparación flexible.
            if not matches:
                matches = (
                    normalized_products_by_model.get(
                        normalized,
                        [],
                    )
                )


            if not matches:
                matches = (
                    normalized_fallback_by_name.get(
                        normalized,
                        [],
                    )
                )


            unique_matches = {
                product.pk: product
                for product in matches
            }


            if not unique_matches:
                result["status"] = (
                    "SIN_PRODUCTO"
                )

                result["detail"] = (
                    "El modelo del Excel no "
                    "coincide con PostgreSQL."
                )

                rows.append(result)
                continue


            if len(unique_matches) > 1:
                result["status"] = (
                    "PRODUCTO_AMBIGUO"
                )

                result["detail"] = (
                    "Coincide con varios "
                    "productos: "
                    + ", ".join(
                        str(pk)
                        for pk
                        in unique_matches
                    )
                )

                rows.append(result)
                continue


            product = next(
                iter(
                    unique_matches.values()
                )
            )


            result["product_id"] = (
                product.pk
            )

            result["product_name"] = (
                product.name
            )


            if (
                product.model_3d
                and product.model_3d.name
                and not options[
                    "replace_existing"
                ]
            ):
                result["status"] = (
                    "YA_TIENE_MODELO"
                )

                result["detail"] = (
                    product.model_3d.name
                )

                rows.append(result)
                continue


            try:
                data, resolved_source = (
                    self.read_source(
                        source=source,
                        excel_dir=excel.parent,
                        models_dir=models_dir,
                        max_bytes=max_bytes,
                    )
                )

                info = validate_glb(
                    data
                )

            except Exception as exc:
                result["status"] = (
                    "FUENTE_INVALIDA"
                )

                result["detail"] = str(
                    exc
                )

                rows.append(result)
                continue


            digest = (
                hashlib.sha256(
                    data
                ).hexdigest()
            )


            result[
                "resolved_source"
            ] = resolved_source

            result["bytes"] = (
                info["bytes"]
            )

            result["sha256"] = (
                digest
            )

            result["status"] = (
                "LISTO_REEMPLAZO"
                if (
                    product.model_3d
                    and product.model_3d.name
                )
                else "LISTO"
            )

            result["detail"] = (
                "GLB 2.0 válido"
            )

            rows.append(result)


        product_ids = [
            item["product_id"]
            for item in rows
            if item["product_id"]
        ]


        duplicates = {
            product_id
            for product_id in product_ids
            if product_ids.count(
                product_id
            ) > 1
        }


        if duplicates:
            for result in rows:
                if (
                    result["product_id"]
                    in duplicates
                ):
                    result["status"] = (
                        "DUPLICADO_EXCEL"
                    )

                    result["detail"] = (
                        "El mismo producto aparece "
                        "más de una vez en el Excel."
                    )


        report = excel.with_name(
            "Modelos3D_reporte.csv"
        )


        self.write_report(
            report,
            rows,
        )


        self.print_rows(
            rows
        )


        errors = [
            row
            for row in rows
            if row["status"]
            in ERROR_STATUSES
        ]


        ready = [
            row
            for row in rows
            if row["status"]
            in READY_STATUSES
        ]


        already = [
            row
            for row in rows
            if row["status"]
            == "YA_TIENE_MODELO"
        ]


        self.stdout.write("")
        self.stdout.write("=" * 78)
        self.stdout.write("RESUMEN")
        self.stdout.write("=" * 78)

        self.stdout.write(
            "Registros técnicos DB : "
            f"{len(technical_records)}"
        )

        self.stdout.write(
            "Filas procesadas      : "
            f"{len(rows)}"
        )

        self.stdout.write(
            "GLB listos            : "
            f"{len(ready)}"
        )

        self.stdout.write(
            "Ya tenían modelo      : "
            f"{len(already)}"
        )

        self.stdout.write(
            "Errores               : "
            f"{len(errors)}"
        )

        self.stdout.write(
            f"Reporte               : "
            f"{report}"
        )

        self.stdout.write("")


        if not apply_changes:
            self.stdout.write(
                self.style.WARNING(
                    "DRY-RUN: no se modificó "
                    "PostgreSQL ni el storage."
                )
            )

            return


        if errors:
            raise CommandError(
                "No se aplicó ningún cambio "
                "porque existen errores. "
                "Corrige primero el reporte."
            )


        if not ready:
            self.stdout.write(
                self.style.WARNING(
                    "No hay modelos nuevos "
                    "para aplicar."
                )
            )

            return


        self.stdout.write("")
        self.stdout.write(
            "Guardando modelos 3D..."
        )


        for index, result in enumerate(
            ready,
            start=1,
        ):
            product = (
                CatalogItem.objects.get(
                    pk=result[
                        "product_id"
                    ]
                )
            )

            data, _ = self.read_source(
                source=result["source"],
                excel_dir=excel.parent,
                models_dir=models_dir,
                max_bytes=max_bytes,
            )

            validate_glb(
                data
            )


            digest = (
                hashlib.sha256(
                    data
                ).hexdigest()
            )


            if (
                digest
                != result["sha256"]
            ):
                raise CommandError(
                    "El archivo cambió después "
                    "de la validación: "
                    f"{result['model']}"
                )


            previous_name = (
                product.model_3d.name
                if (
                    product.model_3d
                    and product.model_3d.name
                )
                else ""
            )


            safe_model = (
                slugify(
                    result["model"]
                )
                or f"producto-{product.pk}"
            )


            filename = (
                f"{safe_model}-"
                f"{digest[:12]}.glb"
            )


            new_name = ""

            try:
                product.model_3d.save(
                    filename,
                    ContentFile(data),
                    save=False,
                )

                new_name = (
                    product.model_3d.name
                )

                product.save(
                    update_fields=[
                        "model_3d"
                    ]
                )

            except Exception:
                if (
                    new_name
                    and storage.exists(
                        new_name
                    )
                ):
                    storage.delete(
                        new_name
                    )

                raise


            if (
                options[
                    "replace_existing"
                ]
                and previous_name
                and previous_name
                != new_name
            ):
                used_elsewhere = (
                    CatalogItem.objects
                    .exclude(
                        pk=product.pk
                    )
                    .filter(
                        model_3d=previous_name
                    )
                    .exists()
                )

                if not used_elsewhere:
                    try:
                        if storage.exists(
                            previous_name
                        ):
                            storage.delete(
                                previous_name
                            )
                    except Exception as exc:
                        self.stderr.write(
                            self.style.WARNING(
                                "No pude eliminar "
                                f"{previous_name}: "
                                f"{exc}"
                            )
                        )


            result["stored_name"] = (
                new_name
            )

            result["status"] = (
                "APLICADO"
            )

            self.stdout.write(
                self.style.SUCCESS(
                    f"[{index}/{len(ready)}] "
                    f"{result['model']} -> "
                    f"{new_name}"
                )
            )


        self.write_report(
            report,
            rows,
        )

        self.stdout.write("")

        self.stdout.write(
            self.style.SUCCESS(
                "IMPORTACIÓN COMPLETADA."
            )
        )

        self.stdout.write(
            f"Reporte final: {report}"
        )


    def read_source(
        self,
        source,
        excel_dir,
        models_dir,
        max_bytes,
    ):
        source = str(
            source
        ).strip()


        if is_url(source):
            request = Request(
                source,
                headers={
                    "User-Agent": (
                        "SMAV-INAHER-3D-Importer/1.0"
                    ),
                    "Accept": (
                        "model/gltf-binary,"
                        "application/octet-stream,"
                        "*/*"
                    ),
                },
            )

            with urlopen(
                request,
                timeout=90,
            ) as response:
                data = response.read(
                    max_bytes + 1
                )

                final_url = (
                    response.geturl()
                )


            if len(data) > max_bytes:
                raise ValueError(
                    "El archivo remoto supera "
                    "el tamaño permitido."
                )


            return data, final_url


        candidates = []


        direct = Path(
            source
        ).expanduser()


        if direct.is_absolute():
            candidates.append(
                direct
            )


        candidates.append(
            excel_dir / source
        )


        if models_dir:
            candidates.append(
                models_dir / source
            )


        basename = (
            filename_from_source(
                source
            )
        )


        if basename:
            candidates.append(
                excel_dir / basename
            )

            if models_dir:
                candidates.append(
                    models_dir / basename
                )


        checked = set()


        for candidate in candidates:
            try:
                candidate = (
                    candidate.resolve()
                )
            except Exception:
                pass


            key = str(
                candidate
            )

            if key in checked:
                continue

            checked.add(key)


            if (
                candidate.exists()
                and candidate.is_file()
            ):
                size = (
                    candidate.stat().st_size
                )

                if size > max_bytes:
                    raise ValueError(
                        "El GLB supera el "
                        "tamaño permitido: "
                        f"{candidate}"
                    )

                return (
                    candidate.read_bytes(),
                    str(candidate),
                )


        raise ValueError(
            "No encontré el archivo 3D. "
            f"Fuente indicada: {source}"
        )


    def write_report(
        self,
        path,
        rows,
    ):
        fields = [
            "excel_row",
            "model",
            "source",
            "product_id",
            "product_name",
            "resolved_source",
            "status",
            "detail",
            "bytes",
            "sha256",
            "stored_name",
        ]


        with path.open(
            "w",
            newline="",
            encoding="utf-8-sig",
        ) as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=fields,
            )

            writer.writeheader()

            for row in rows:
                writer.writerow(
                    {
                        field:
                            row.get(
                                field,
                                "",
                            )
                        for field
                        in fields
                    }
                )


    def print_rows(
        self,
        rows,
    ):
        for result in rows:
            status = result[
                "status"
            ]

            model = result[
                "model"
            ]

            product_name = (
                result[
                    "product_name"
                ]
                or "-"
            )


            if status in ERROR_STATUSES:
                output = (
                    self.style.ERROR(
                        status
                    )
                )
            elif status in READY_STATUSES:
                output = (
                    self.style.SUCCESS(
                        status
                    )
                )
            else:
                output = (
                    self.style.WARNING(
                        status
                    )
                )


            self.stdout.write(
                f"Fila "
                f"{result['excel_row']:>3} | "
                f"{model:<22} | "
                f"{product_name:<28} | "
                f"{output}"
            )


            if (
                status
                in ERROR_STATUSES
            ):
                self.stdout.write(
                    "       -> "
                    + result[
                        "detail"
                    ]
                )
